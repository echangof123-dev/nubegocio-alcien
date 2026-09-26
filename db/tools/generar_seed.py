#!/usr/bin/env python3
"""Genera db/seed/catalogo.sql a partir de data/catalogo_plantillas_negocios.xlsx.

El Excel es la fuente del catálogo (módulos, familias, matriz, tipos de negocio y fichas).
Aquí se agregan los datos que no viven en el Excel: fase de cada módulo, dependencias
entre módulos y los planes con sus módulos.

Uso:  python3 db/tools/generar_seed.py
"""
from __future__ import annotations

import os
import sys
from openpyxl import load_workbook

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
EXCEL = os.path.join(RAIZ, "data", "catalogo_plantillas_negocios.xlsx")
SALIDA = os.path.join(RAIZ, "db", "seed", "catalogo.sql")

# Fase en la que se construye cada módulo (1 = MVP).
FASE = {
    **{m: 1 for m in ["M01", "M02", "M03", "M04", "M05", "M06", "M07", "M14", "M15", "M19", "M20", "M21"]},
    **{m: 2 for m in ["M08", "M09", "M10", "M11", "M12", "M13", "M16", "M17", "M23", "M24"]},
    **{m: 3 for m in ["M18", "M22", "M25", "M26"]},
}

# módulo → módulos que necesita para funcionar
DEPENDENCIAS = {
    "M05": ["M04"],  # variantes → inventario
    "M06": ["M04"],  # venta por peso → inventario
    "M07": ["M04"],  # código de barras → inventario
    "M08": ["M04"],  # recetas → inventario
    "M15": ["M04"],  # proveedores y compras → inventario
    "M16": ["M04"],  # lotes → inventario
    "M23": ["M04"],  # series → inventario
    "M25": ["M04", "M15"],  # acopio → inventario, compras
    "M09": ["M08"],  # mesas y comandas → recetas
}

PLANES = [
    # codigo, nombre, precio, orden, limites
    ("gratis", "Gratis", 0.00, 1, '{"usuarios": 1, "productos": 50}'),
    ("emprendedor", "Emprendedor", 7.99, 2, '{"usuarios": 2}'),
    ("negocio", "Negocio", 17.99, 3, '{"usuarios": 5}'),
    ("pro", "Pro", 29.99, 4, '{"usuarios": 15, "sucursales": 3}'),
]
PLAN_GRATIS = ["M01", "M02", "M03", "M14", "M21"]
PLAN_EMPRENDEDOR = PLAN_GRATIS + ["M04", "M05", "M06", "M07"]
PLAN_NEGOCIO = PLAN_EMPRENDEDOR + ["M15", "M19", "M20"]
# Pro incluye todos los módulos (se calcula con los del Excel).

# Tipos más comunes: ganan los empates en la búsqueda ("tienda" → Tienda de barrio primero).
PRIORIDAD = {
    "Tienda de barrio": 5, "Minimarket": 4, "Restaurante": 4, "Comida rápida": 4, "Tienda de ropa": 4,
    "Peluquería": 4, "Ferretería": 4, "Farmacia": 4, "Almacén agrícola / agroquímicos": 4,
    "Cafetería": 3, "Panadería": 3, "Bazar": 3, "Barbería": 3, "Taller mecánico": 3,
    "Agroveterinaria": 3, "Tienda de celulares": 3, "Zapatería": 3, "Papelería y útiles escolares": 3,
}

# Sinónimos extra, además de los del Excel, para frases comunes al registrarse.
SINONIMOS_EXTRA = {
    "Tienda de barrio": ["tiendita", "abarrotes", "víveres", "despensa"],
    "Tienda de ropa": ["ropa", "prendas", "vestir"],
    "Almacén agrícola / agroquímicos": ["insumos para el campo", "agropecuario", "fertilizantes", "finca", "agro"],
    "Comida rápida": ["comidas", "picadas"],
    "Restaurante": ["comida", "almuerzos", "restaurant"],
}


def q(v) -> str:
    """Literal SQL seguro para texto (o NULL)."""
    if v is None:
        return "null"
    s = str(v).strip()
    if s == "":
        return "null"
    return "'" + s.replace("'", "''") + "'"


def arr(items: list[str]) -> str:
    if not items:
        return "'{}'::text[]"
    return "array[" + ", ".join(q(i) for i in items) + "]::text[]"


def partir(texto, sep) -> list[str]:
    if not texto:
        return []
    return [p.strip() for p in str(texto).split(sep) if p.strip()]


def filas(ws, desde=2):
    for fila in ws.iter_rows(min_row=desde, values_only=True):
        if fila and fila[0] is not None and str(fila[0]).strip():
            yield fila


def main() -> int:
    wb = load_workbook(EXCEL, data_only=False)

    # --- Módulos ---
    modulos = []
    for f in filas(wb["Módulos"]):
        codigo, nombre, desc, tipo = f[:4]
        if not str(codigo).startswith("M"):
            continue
        modulos.append((codigo, nombre, desc, "nucleo" if str(tipo).lower().startswith("n") else "activable"))
    codigos_modulo = {m[0] for m in modulos}
    faltan = codigos_modulo - FASE.keys()
    if faltan:
        print(f"Módulos sin fase asignada: {sorted(faltan)}", file=sys.stderr)
        return 1

    # --- Familias ---
    familias = []
    for f in filas(wb["Familias"]):
        codigo, nombre, desc, palabra, unidad = f[:5]
        if str(codigo).startswith("F"):
            familias.append((codigo, nombre, desc, palabra, unidad))

    # --- Matriz familia × módulo ---
    wsm = wb["Matriz"]
    encabezado = [c.value for c in wsm[1]]
    columnas = {}
    for i, h in enumerate(encabezado):
        if h and str(h).startswith("M") and str(h)[:3] in codigos_modulo:
            columnas[i] = str(h)[:3]
    matriz = []
    for f in wsm.iter_rows(min_row=2, values_only=True):
        if not f[0] or not str(f[0]).startswith("F"):
            continue
        for i, mod in columnas.items():
            v = (f[i] or "").strip() if isinstance(f[i], str) else f[i]
            if v in ("S", "O"):
                matriz.append((f[0], mod, v))

    # --- Tipos de negocio ---
    tipos = []
    for f in filas(wb["Tipos de negocio"]):
        codigo, nombre, fam, _fam_nombre, sinon, cats, palabra, unidad, ajuste = f[:9]
        if not str(codigo).startswith("T"):
            continue
        tipos.append({
            "codigo": codigo, "nombre": nombre, "familia": fam,
            "sinonimos": partir(sinon, ","), "categorias": partir(cats, ";"),
            "palabra": palabra, "unidad": unidad, "ajuste": ajuste,
            "prioridad": PRIORIDAD.get(nombre, 0),
        })
        tipos[-1]["sinonimos"] += [s for s in SINONIMOS_EXTRA.get(nombre, []) if s not in tipos[-1]["sinonimos"]]
    por_nombre = {t["nombre"]: t["codigo"] for t in tipos}
    for n in list(PRIORIDAD) + list(SINONIMOS_EXTRA):
        if n not in por_nombre:
            print(f"Tipo desconocido en PRIORIDAD/SINONIMOS_EXTRA: {n!r}", file=sys.stderr)
            return 1

    # --- Fichas de productos ---
    fichas = []
    orden = {}
    for f in filas(wb["Fichas ejemplo"]):
        tipo, cat, prod, unidad = f[:4]
        if not cat or not prod:
            continue  # notas al pie de la hoja
        if tipo not in por_nombre:
            print(f"Ficha de un tipo que no existe: {tipo!r}", file=sys.stderr)
            return 1
        unidad = str(unidad or "Unidad")
        variantes = None
        if "talla" in unidad.lower() or "·" in unidad:
            variantes, unidad = unidad, "Unidad"
        orden[tipo] = orden.get(tipo, 0) + 1
        fichas.append((por_nombre[tipo], cat, prod, unidad, variantes, orden[tipo]))

    # --- SQL ---
    out = [
        "-- GENERADO por db/tools/generar_seed.py desde data/catalogo_plantillas_negocios.xlsx",
        "-- No editar a mano: cambia el Excel o el script y vuelve a generar.",
        "",
        "insert into catalogo.plan (codigo, nombre, precio_mensual, orden, limites) values",
        ",\n".join(f"  ({q(c)}, {q(n)}, {p:.2f}, {o}, {q(l)}::jsonb)" for c, n, p, o, l in PLANES)
        + "\non conflict (codigo) do update set nombre = excluded.nombre, precio_mensual = excluded.precio_mensual,"
          " orden = excluded.orden, limites = excluded.limites;",
        "",
        "insert into catalogo.modulo (codigo, nombre, descripcion, tipo, fase) values",
        ",\n".join(f"  ({q(c)}, {q(n)}, {q(d)}, {q(t)}, {FASE[c]})" for c, n, d, t in modulos)
        + "\non conflict (codigo) do update set nombre = excluded.nombre, descripcion = excluded.descripcion,"
          " tipo = excluded.tipo, fase = excluded.fase;",
        "",
        "insert into catalogo.modulo_dependencia (modulo, requiere) values",
        ",\n".join(f"  ({q(m)}, {q(r)})" for m, reqs in DEPENDENCIAS.items() for r in reqs)
        + "\non conflict do nothing;",
        "",
    ]

    plan_mods = {
        "gratis": PLAN_GRATIS, "emprendedor": PLAN_EMPRENDEDOR,
        "negocio": PLAN_NEGOCIO, "pro": sorted(codigos_modulo),
    }
    out.append("insert into catalogo.plan_modulo (plan, modulo) values")
    out.append(",\n".join(f"  ({q(p)}, {q(m)})" for p, ms in plan_mods.items() for m in ms)
               + "\non conflict do nothing;")
    out.append("")

    out.append("insert into catalogo.familia (codigo, nombre, descripcion, palabra_items, unidad_defecto) values")
    out.append(",\n".join(f"  ({q(c)}, {q(n)}, {q(d)}, {q(p)}, {q(u)})" for c, n, d, p, u in familias)
               + "\non conflict (codigo) do update set nombre = excluded.nombre, descripcion = excluded.descripcion,"
                 " palabra_items = excluded.palabra_items, unidad_defecto = excluded.unidad_defecto,"
                 " version = catalogo.familia.version + 1;")
    out.append("")

    out.append("insert into catalogo.familia_modulo (familia, modulo, modo) values")
    out.append(",\n".join(f"  ({q(f)}, {q(m)}, {q(v)})" for f, m, v in matriz)
               + "\non conflict (familia, modulo) do update set modo = excluded.modo;")
    out.append("")

    out.append("insert into catalogo.tipo_negocio (codigo, nombre, familia, sinonimos, categorias, palabra_items, unidad, ajuste, prioridad) values")
    out.append(",\n".join(
        f"  ({q(t['codigo'])}, {q(t['nombre'])}, {q(t['familia'])}, {arr(t['sinonimos'])}, {arr(t['categorias'])}, "
        f"{q(t['palabra'])}, {q(t['unidad'])}, {q(t['ajuste'])}, {t['prioridad']})" for t in tipos)
        + "\non conflict (codigo) do update set nombre = excluded.nombre, familia = excluded.familia,"
          " sinonimos = excluded.sinonimos, categorias = excluded.categorias, palabra_items = excluded.palabra_items,"
          " unidad = excluded.unidad, ajuste = excluded.ajuste, prioridad = excluded.prioridad,"
          " version = catalogo.tipo_negocio.version + 1;")
    out.append("")

    out.append("insert into catalogo.plantilla_producto (tipo_negocio_id, categoria, nombre, unidad, variantes, orden)")
    out.append("select tn.id, v.categoria, v.nombre, v.unidad, v.variantes, v.orden from (values")
    out.append(",\n".join(f"  ({q(c)}, {q(cat)}, {q(p)}, {q(u)}, {q(var)}, {o})" for c, cat, p, u, var, o in fichas))
    out.append(") as v(codigo, categoria, nombre, unidad, variantes, orden)")
    out.append("join catalogo.tipo_negocio tn on tn.codigo = v.codigo")
    out.append("on conflict (tipo_negocio_id, nombre) do update set categoria = excluded.categoria,"
               " unidad = excluded.unidad, variantes = excluded.variantes, orden = excluded.orden;")
    out.append("")

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))
    print(f"{SALIDA}: {len(modulos)} módulos, {len(familias)} familias, {len(matriz)} celdas de matriz, "
          f"{len(tipos)} tipos, {len(fichas)} productos de ejemplo, {len(PLANES)} planes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
